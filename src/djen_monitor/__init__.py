from __future__ import annotations

from openpyxl.worksheet.worksheet import Worksheet

from .constants import APP_VERSION


def _apply_excel_desktop_compatibility() -> None:
    """Evita tabelas estruturadas que algumas versões do Excel reparavam ao abrir.

    O relatório já usa AutoFilter da própria planilha e formatação manual. As
    estruturas Table do openpyxl eram redundantes e geravam /xl/tables/table*.xml,
    que o Excel desktop podia remover durante a reparação do arquivo. Mantemos o
    filtro normal e suprimimos apenas o registro dessas tabelas estruturadas.
    """
    if getattr(Worksheet, "_djen_original_add_table", None) is None:
        Worksheet._djen_original_add_table = Worksheet.add_table  # type: ignore[attr-defined]

    def _ignore_structured_table(self: Worksheet, table) -> None:
        return None

    Worksheet.add_table = _ignore_structured_table  # type: ignore[method-assign]


_apply_excel_desktop_compatibility()

__version__ = APP_VERSION
